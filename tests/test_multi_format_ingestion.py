"""Tests for Sprint 5 — multi-format ingestion (TXT, MD, DOCX, EPUB) + detect."""

from __future__ import annotations

from pathlib import Path

import pytest

from core.config_loader import Config
from core.exceptions import ConfigValidationError, NCNError
from ingestion.detect import (
    detect_source_type,
    is_url,
    supported_extensions,
    supported_source_types,
)
from ingestion.docx_ingestor import DocxIngestor
from ingestion.epub_ingestor import EpubIngestor
from ingestion.factory import create_ingestor, supported_ingestor_types
from ingestion.markdown_ingestor import MarkdownIngestor
from ingestion.txt_ingestor import TXTIngestor


# ---------------------------------------------------------------------------
# detect_source_type
# ---------------------------------------------------------------------------

class TestDetectSourceType:
    """Source type detection by extension and Content-Type."""

    @pytest.mark.parametrize("path,expected", [
        ("manual.pdf", "pdf"),
        ("guide.PDF", "pdf"),
        ("book.epub", "epub"),
        ("doc.docx", "docx"),
        ("legacy.doc", "docx"),
        ("readme.md", "markdown"),
        ("notes.markdown", "markdown"),
        ("log.txt", "txt"),
        ("page.html", "html"),
        ("page.htm", "html"),
        ("unknown.xyz", "text"),
        ("no_extension", "text"),
    ])
    def test_extension_detection(self, path, expected):
        assert detect_source_type(path) == expected

    @pytest.mark.parametrize("content_type,expected", [
        ("application/pdf", "pdf"),
        ("application/pdf; charset=binary", "pdf"),
        ("application/epub+zip", "epub"),
        ("text/html; charset=utf-8", "html"),
        ("text/markdown", "markdown"),
        ("text/plain", "txt"),
        ("application/msword", "docx"),
    ])
    def test_content_type_detection(self, content_type, expected):
        # Use an extensionless URL so only content-type decides
        assert detect_source_type("https://ex.com/resource", content_type) == expected

    def test_content_type_overrides_extension(self):
        """A PDF served from a .html URL should be detected as PDF."""
        assert detect_source_type("https://ex.com/doc.html", "application/pdf") == "pdf"

    def test_url_with_query_params(self):
        """Query params should not confuse extension detection."""
        assert detect_source_type("https://ex.com/book.epub?version=2") == "epub"

    def test_is_url(self):
        assert is_url("https://example.com/a.pdf")
        assert is_url("http://example.com")
        assert not is_url("/local/file.pdf")
        assert not is_url("file.pdf")
        assert not is_url("not a url at all")

    def test_supported_lists(self):
        exts = supported_extensions()
        types = supported_source_types()
        assert ".pdf" in exts
        assert ".epub" in exts
        assert ".md" in exts
        assert "pdf" in types
        assert "epub" in types


# ---------------------------------------------------------------------------
# Factory
# ---------------------------------------------------------------------------

class TestIngestorFactory:
    """Factory returns the right ingestor class per source_type."""

    def test_supported_types(self):
        types = supported_ingestor_types()
        assert "text" in types
        assert "txt" in types
        assert "pdf" in types
        assert "docx" in types
        assert "epub" in types
        assert "markdown" in types
        assert "md" in types

    def test_creates_txt_ingestor(self):
        ing = create_ingestor("txt", Config())
        assert isinstance(ing, TXTIngestor)

    def test_creates_markdown_ingestor(self):
        ing = create_ingestor("markdown", Config())
        assert isinstance(ing, MarkdownIngestor)

    def test_md_alias(self):
        ing = create_ingestor("md", Config())
        assert isinstance(ing, MarkdownIngestor)

    def test_creates_docx_ingestor(self):
        ing = create_ingestor("docx", Config())
        assert isinstance(ing, DocxIngestor)

    def test_creates_epub_ingestor(self):
        ing = create_ingestor("epub", Config())
        assert isinstance(ing, EpubIngestor)

    def test_unknown_type_raises(self):
        with pytest.raises(ConfigValidationError):
            create_ingestor("unknown_format", Config())


# ---------------------------------------------------------------------------
# TXTIngestor
# ---------------------------------------------------------------------------

class TestTXTIngestor:

    async def test_ingest_short_text_file(self, tmp_path):
        path = tmp_path / "hello.txt"
        path.write_text("Hello world.\n\nThis is a test.\n", encoding="utf-8")

        ing = TXTIngestor(Config())
        chunks = await ing.ingest(str(path))

        assert len(chunks) >= 1
        assert all(c.source_type == "txt" for c in chunks)
        assert any("Hello world" in c.text for c in chunks)

    async def test_ingest_latin1_encoded_file(self, tmp_path):
        path = tmp_path / "latin.txt"
        path.write_bytes("Acción y café en español.".encode("latin-1"))

        ing = TXTIngestor(Config())
        chunks = await ing.ingest(str(path))
        assert len(chunks) >= 1
        # Decoded content should be readable (may have replacement chars but not crash)
        assert any(len(c.text) > 5 for c in chunks)

    async def test_missing_file_raises(self, tmp_path):
        ing = TXTIngestor(Config())
        with pytest.raises(NCNError):
            await ing.ingest(str(tmp_path / "nope.txt"))

    async def test_empty_file_returns_no_chunks(self, tmp_path):
        path = tmp_path / "empty.txt"
        path.write_text("", encoding="utf-8")
        ing = TXTIngestor(Config())
        chunks = await ing.ingest(str(path))
        assert chunks == []

    async def test_long_file_gets_chunked(self, tmp_path):
        # 6000 chars of text, should split into multiple chunks
        long_text = "\n\n".join([f"Paragraph {i} " + "x" * 100 for i in range(50)])
        path = tmp_path / "long.txt"
        path.write_text(long_text, encoding="utf-8")

        ing = TXTIngestor(Config())
        chunks = await ing.ingest(str(path))
        assert len(chunks) > 1


# ---------------------------------------------------------------------------
# MarkdownIngestor
# ---------------------------------------------------------------------------

class TestMarkdownIngestor:

    async def test_ingest_with_headers(self, tmp_path):
        md = (
            "# Survival Guide\n\n"
            "Intro text.\n\n"
            "## Water\n\n"
            "How to purify water: boil for 1 minute.\n\n"
            "## Food\n\n"
            "Identify safe plants.\n"
        )
        path = tmp_path / "guide.md"
        path.write_text(md, encoding="utf-8")

        ing = MarkdownIngestor(Config())
        chunks = await ing.ingest(str(path))

        assert len(chunks) >= 1
        assert all(c.source_type == "markdown" for c in chunks)
        combined = "\n".join(c.text for c in chunks)
        assert "Water" in combined and "Food" in combined

    async def test_ingest_without_headers_falls_back(self, tmp_path):
        md = "\n\n".join(["Paragraph " + str(i) for i in range(5)])
        path = tmp_path / "no_headers.md"
        path.write_text(md, encoding="utf-8")

        ing = MarkdownIngestor(Config())
        chunks = await ing.ingest(str(path))
        assert len(chunks) >= 1

    async def test_strips_markdown_formatting(self, tmp_path):
        md = "# Title\n\nSome **bold** and *italic* text with `code`.\n"
        path = tmp_path / "fmt.md"
        path.write_text(md, encoding="utf-8")

        ing = MarkdownIngestor(Config())
        chunks = await ing.ingest(str(path))
        combined = "\n".join(c.text for c in chunks)
        # The content is preserved, but raw markdown marks should not leak
        assert "bold" in combined and "italic" in combined
        assert "**bold**" not in combined


# ---------------------------------------------------------------------------
# DocxIngestor
# ---------------------------------------------------------------------------

class TestDocxIngestor:

    async def test_ingest_simple_docx(self, tmp_path):
        from docx import Document

        path = tmp_path / "doc.docx"
        doc = Document()
        doc.add_heading("Survival Protocol", level=1)
        doc.add_paragraph("First aid is critical in emergencies.")
        doc.add_paragraph("Always check for breathing first.")
        doc.save(str(path))

        ing = DocxIngestor(Config())
        chunks = await ing.ingest(str(path))

        assert len(chunks) >= 1
        assert all(c.source_type == "docx" for c in chunks)
        combined = "\n".join(c.text for c in chunks)
        assert "breathing" in combined

    async def test_ingest_docx_with_table(self, tmp_path):
        from docx import Document

        path = tmp_path / "table.docx"
        doc = Document()
        doc.add_paragraph("Contacts")
        table = doc.add_table(rows=2, cols=2)
        table.rows[0].cells[0].text = "Name"
        table.rows[0].cells[1].text = "Role"
        table.rows[1].cells[0].text = "Alice"
        table.rows[1].cells[1].text = "Doctor"
        doc.save(str(path))

        ing = DocxIngestor(Config())
        chunks = await ing.ingest(str(path))
        combined = "\n".join(c.text for c in chunks)
        assert "Alice" in combined and "Doctor" in combined
        # Cells should be pipe-separated
        assert "|" in combined

    async def test_missing_docx_raises(self, tmp_path):
        ing = DocxIngestor(Config())
        with pytest.raises(NCNError):
            await ing.ingest(str(tmp_path / "missing.docx"))


# ---------------------------------------------------------------------------
# EpubIngestor
# ---------------------------------------------------------------------------

class TestEpubIngestor:

    async def test_ingest_simple_epub(self, tmp_path):
        from ebooklib import epub

        path = tmp_path / "book.epub"
        book = epub.EpubBook()
        book.set_identifier("id-123")
        book.set_title("Survival Manual")
        book.set_language("en")

        chapter1 = epub.EpubHtml(title="Chapter 1", file_name="ch1.xhtml", lang="en")
        chapter1.content = (
            "<html><body><h1>Chapter 1</h1>"
            "<p>Water purification is essential for survival.</p>"
            "<p>Boiling water for one minute kills most pathogens.</p>"
            "</body></html>"
        )
        chapter2 = epub.EpubHtml(title="Chapter 2", file_name="ch2.xhtml", lang="en")
        chapter2.content = (
            "<html><body><h1>Chapter 2</h1>"
            "<p>Fire-starting techniques using flint and tinder.</p>"
            "</body></html>"
        )
        book.add_item(chapter1)
        book.add_item(chapter2)
        book.toc = (chapter1, chapter2)
        book.add_item(epub.EpubNcx())
        book.add_item(epub.EpubNav())
        book.spine = ["nav", chapter1, chapter2]

        epub.write_epub(str(path), book, {})

        ing = EpubIngestor(Config())
        chunks = await ing.ingest(str(path))

        assert len(chunks) >= 1
        assert all(c.source_type == "epub" for c in chunks)
        combined = "\n".join(c.text for c in chunks)
        assert "Water purification" in combined
        assert "Fire-starting" in combined

    async def test_missing_epub_raises(self, tmp_path):
        ing = EpubIngestor(Config())
        with pytest.raises(NCNError):
            await ing.ingest(str(tmp_path / "missing.epub"))
