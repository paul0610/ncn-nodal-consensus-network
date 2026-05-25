"""DocxIngestor — ingest Microsoft Word .docx files via python-docx.

Extracts paragraph text and table contents. Tables are flattened into
plain text (cell1 | cell2 | cell3) so the consensus can reason over them
without knowing about table structure.

Note: legacy .doc files are not supported natively (would need antiword or
similar). If passed a .doc, this ingestor attempts python-docx anyway;
failures are surfaced as NCNError.
"""

from __future__ import annotations

import asyncio
from pathlib import Path

from loguru import logger

from core.config_loader import Config
from core.exceptions import NCNError
from core.models import RawContent
from ingestion.base import IngestorPort


_MAX_FILE_SIZE_MB = 50
_CHUNK_TARGET_CHARS = 2000


class DocxIngestor(IngestorPort):
    """Extract paragraph and table text from a .docx file."""

    def __init__(self, config: Config) -> None:
        self.config = config

    async def ingest(self, source: str) -> list[RawContent]:
        path = Path(source)
        if not path.exists():
            raise NCNError(f"DOCX file not found: {source}")

        size_mb = path.stat().st_size / (1024 * 1024)
        if size_mb > _MAX_FILE_SIZE_MB:
            raise NCNError(
                f"DOCX file too large ({size_mb:.1f} MB > {_MAX_FILE_SIZE_MB} MB)"
            )

        return await asyncio.to_thread(self._extract, str(path))

    @staticmethod
    def _extract(docx_path: str) -> list[RawContent]:
        try:
            from docx import Document
        except ImportError as exc:
            raise NCNError(
                "python-docx not installed. Run: pip install python-docx"
            ) from exc

        try:
            doc = Document(docx_path)
        except Exception as exc:
            raise NCNError(f"Cannot open DOCX '{docx_path}': {exc}") from exc

        # Collect all paragraph text
        blocks: list[str] = []

        for para in doc.paragraphs:
            text = para.text.strip()
            if text:
                blocks.append(text)

        # Flatten tables into pipe-separated rows
        for table in doc.tables:
            for row in table.rows:
                cells = [cell.text.strip() for cell in row.cells]
                cells = [c for c in cells if c]
                if cells:
                    blocks.append(" | ".join(cells))

        # Chunk the collected blocks to the target size
        name = Path(docx_path).name
        chunks = _merge_to_chunks(blocks, _CHUNK_TARGET_CHARS)

        contents = [
            RawContent(
                text=chunk,
                source=f"docx:{name}:chunk{i + 1}",
                source_type="docx",
            )
            for i, chunk in enumerate(chunks)
            if chunk.strip()
        ]
        logger.info(f"DOCX ingested: {docx_path} → {len(contents)} chunks")
        return contents


def _merge_to_chunks(blocks: list[str], target_chars: int) -> list[str]:
    """Combine paragraph/row blocks into chunks near *target_chars*."""
    chunks: list[str] = []
    current = ""
    for block in blocks:
        if len(current) + len(block) + 2 > target_chars and current:
            chunks.append(current.strip())
            current = block
        else:
            current = f"{current}\n\n{block}" if current else block
    if current.strip():
        chunks.append(current.strip())
    return chunks
